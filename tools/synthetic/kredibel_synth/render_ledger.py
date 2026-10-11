"""The LedgerLite 3 export (docs/18 §3.5).

A CSV with debit and credit columns, and an XLSX variant with a different header and one signed
amount column, to test a second mapping.
"""

import csv
import io
import re
import zipfile
from datetime import datetime

from openpyxl import Workbook

from .model import LedgerRow
from .world import BUSINESS, WATERMARK

CSV_HEADER = [
    "Date",
    "Journal",
    "Account",
    "Account Name",
    "Description",
    "Reference",
    "Debit",
    "Credit",
]
XLSX_HEADER = ["Posting date", "Acct", "Acct name", "Narration", "Ref", "Amount"]
FIXED_TIME = datetime(2026, 7, 1, 9, 0, 0)


def comment_line() -> str:
    return f"# {WATERMARK}. LedgerLite 3 export, {BUSINESS.name}, 01/01/2026 to 30/06/2026"


def render_ledger_csv(rows: list[LedgerRow]) -> bytes:
    out = io.StringIO()
    out.write(comment_line() + "\n")
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(CSV_HEADER)
    for r in rows:
        writer.writerow(
            [
                f"{r.date:%d/%m/%Y}",
                r.journal,
                r.account,
                r.account_name,
                r.description,
                r.reference,
                f"{r.debit:.2f}" if r.debit is not None else "",
                f"{r.credit:.2f}" if r.credit is not None else "",
            ]
        )
    return out.getvalue().encode("utf-8")


def render_ledger_xlsx(rows: list[LedgerRow]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Journal"
    ws.append([comment_line()])
    ws.append(XLSX_HEADER)
    for r in rows:
        signed = r.debit if r.debit is not None else -(r.credit or 0)
        ws.append([r.date, r.account, r.account_name, r.description, r.reference, float(signed)])
        ws.cell(row=ws.max_row, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=ws.max_row, column=6).number_format = "#,##0.00"
    wb.properties.creator = "LedgerLite 3"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.properties.lastModifiedBy = "LedgerLite 3"
    raw = io.BytesIO()
    wb.save(raw)
    return _fixed_zip(raw.getvalue())


def _fixed_zip(data: bytes) -> bytes:
    """openpyxl stamps zip entries with the current time; rewrite them with a fixed one."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=FIXED_TIME.timetuple()[:6])
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
                content = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>" + stamp, content
                )
            target.writestr(fixed, content)
    return out.getvalue()
