"""Rendered files (docs/18 §3, §7; AC-002-2…AC-002-6)."""

import csv
import hashlib
import io
from collections import Counter
from decimal import Decimal
from functools import cache

import pdfplumber
import pytest
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from kredibel_synth import world
from kredibel_synth.model import fmt
from kredibel_synth.output import LEDGER_FILE, Bundle, render_books
from kredibel_synth.render_ledger import render_ledger_xlsx
from tests.conftest import books_for


@cache
def bundle(seed: int = 1) -> Bundle:
    return render_books(books_for(seed))


def _text(pdf: bytes) -> list[str]:
    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        return [page.extract_text() or "" for page in doc.pages]


def test_ac002_1_one_file_per_document_and_the_ledger():
    files = bundle().files
    assert len([f for f in files if f.startswith("statement-")]) == 6
    assert len([f for f in files if f.startswith("inv-sales-")]) == 24
    assert len([f for f in files if f.startswith("inv-purchase-")]) == 16
    receipts = [f for f in files if f.startswith("receipt-")]
    assert len(receipts) == 24 and {f.rsplit(".", 1)[1] for f in receipts} == {"jpg", "pdf"}
    assert LEDGER_FILE in files


@pytest.mark.parametrize("month", world.MONTHS)
def test_ac002_4_statements_are_two_or_three_pages_with_fixed_metadata(month: str):
    pdf = bundle().files[f"statement-{month}.pdf"]
    reader = PdfReader(io.BytesIO(pdf))
    assert 2 <= len(reader.pages) <= 3
    meta = reader.metadata
    assert meta is not None
    assert meta["/Producer"] == world.STATEMENT_PRODUCER
    assert meta["/Creator"] == world.STATEMENT_CREATOR
    st = books_for(1).statements[month]
    assert meta["/CreationDate"] == f"D:{st.created_at:%Y%m%d}020000+07'00'"
    assert meta["/ModDate"] == meta["/CreationDate"]
    assert pdf.count(b"%%EOF") == 1  # a single revision


def test_ac002_4_every_row_and_balance_is_on_the_statement_in_english_format():
    for month, st in books_for(1).statements.items():
        text = "\n".join(_text(bundle().files[f"statement-{month}.pdf"]))
        assert "OPENING BALANCE" in text and "CLOSING BALANCE" in text
        assert fmt(st.opening) in text and fmt(st.closing) in text
        for row in st.rows:
            assert fmt(row.balance) in text, row.key
            assert fmt(row.amount) in text, row.key


def test_ac002_4_amounts_use_the_embedded_font():
    pdf = bundle().files["statement-2026-03.pdf"]
    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        fonts = {
            ch["fontname"].split("+")[-1]
            for page in doc.pages
            for ch in page.chars
            if ch["text"].isdigit()
        }
    assert fonts <= {"DejaVuSans", "DejaVuSans-Bold"}


def test_ac002_4_the_sidecar_map_places_every_amount():
    rendered = bundle().statements["2026-03"]
    st = books_for(1).statements["2026-03"]
    for row in st.rows:
        placed = rendered.amounts[row.key]
        fields = {p.field for p in placed}
        assert ("credit" if row.credit is not None else "debit") in fields and "balance" in fields
        assert rendered.row_pages[row.key] in range(1, rendered.pages + 1)


def test_ac002_3_every_pdf_page_carries_the_watermark():
    for name, data in bundle().files.items():
        if name.endswith(".pdf"):
            assert all(world.WATERMARK in page for page in _text(data)), name


def test_ac002_5_invoices_show_their_totals_and_pay_to():
    for inv in books_for(1).invoices:
        text = _text(bundle().files[inv.file])[0]
        assert inv.number in text and fmt(inv.total) in text and fmt(inv.vat) in text
        assert inv.pay_to.account_number in text


def test_ac002_5_each_supplier_has_its_own_font():
    fonts = {}
    for inv in books_for(1).purchases:
        with pdfplumber.open(io.BytesIO(bundle().files[inv.file])) as doc:
            names = Counter(
                ch["fontname"].split("+")[-1].split("-")[0] for ch in doc.pages[0].chars
            )
            fonts[inv.layout] = names.most_common(1)[0][0]
    assert len(set(fonts.values())) == len(fonts)


def test_receipt_photos_are_jpeg_without_exif():
    for rec in books_for(1).receipts:
        data = bundle().files[rec.file]
        if rec.form == "jpeg":
            image = Image.open(io.BytesIO(data))
            assert image.format == "JPEG" and not image.getexif()
        else:
            assert fmt(rec.total) in _text(data)[0]


def test_ac002_3_and_6_the_ledger_csv_opens_with_the_comment_and_mirrors_the_rows():
    lines = bundle().files[LEDGER_FILE].decode().splitlines()
    assert lines[0].startswith(f"# {world.WATERMARK}")
    rows = list(csv.DictReader(lines[1:]))
    ledger = books_for(1).ledger
    assert len(rows) == len(ledger)
    debits = sum(Decimal(r["Debit"] or 0) for r in rows)
    credits = sum(Decimal(r["Credit"] or 0) for r in rows)
    assert debits == credits


def test_the_xlsx_variant_has_one_signed_amount_column():
    wb = load_workbook(io.BytesIO(render_ledger_xlsx(books_for(1).ledger)))
    ws = wb.active
    assert ws is not None
    assert str(ws.cell(1, 1).value).startswith(f"# {world.WATERMARK}")
    assert [c.value for c in ws[2]] == [
        "Posting date",
        "Acct",
        "Acct name",
        "Narration",
        "Ref",
        "Amount",
    ]
    amounts = [Decimal(str(row[5])) for row in ws.iter_rows(min_row=3, values_only=True)]
    assert len(amounts) == len(books_for(1).ledger) and sum(amounts) == 0


def test_ac002_2_rendering_twice_gives_identical_bytes():
    first = render_books(books_for(2))
    second = render_books(books_for(2))
    digest = lambda b: {k: hashlib.sha256(v).hexdigest() for k, v in b.files.items()}  # noqa: E731
    assert digest(first) == digest(second)
    assert render_ledger_xlsx(books_for(2).ledger) == render_ledger_xlsx(books_for(2).ledger)
